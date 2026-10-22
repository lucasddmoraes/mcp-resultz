#!/usr/bin/env python3
"""
Relatório de Vendas WooCommerce - GDV
Monta a planilha .xlsx a partir das respostas da ferramenta relatorio_vendas do conector
Resultz Lojas (JSON) ou de um CSV exportado do WooCommerce. Não acessa a loja e não usa
credenciais: os dados chegam pelo conector ou pelo arquivo do usuário.
"""

import argparse
import csv
import json
import os
import sys
from datetime import datetime

try:
    import openpyxl
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
except ImportError:
    os.system(f"{sys.executable} -m pip install openpyxl --break-system-packages -q")
    import openpyxl
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

COLUNAS = ["Pedido", "Data", "Status", "Método de Pagamento", "Email", "Aluno", "Produto", "Qtde", "Valor"]
LARGURAS = {"Pedido": 12, "Data": 18, "Status": 12, "Método de Pagamento": 25, "Email": 30,
            "Aluno": 22, "Produto": 32, "Qtde": 8, "Valor": 12}
STATUS_PT = {"completed": "Concluído", "processing": "Processando", "on-hold": "Aguardando",
             "pending": "Pendente", "cancelled": "Cancelado", "refunded": "Reembolsado", "failed": "Falhou"}
STATUS_FORA = {"cancelled", "refunded", "failed"}
EMAILS_INTERNOS = {"login@resultzads.com"}


def extrair_nome_aluno(email: str) -> str:
    """rafaella.borelli@gdv.g12.br -> Rafaella Borelli. Outros domínios -> vazio."""
    if not email or not email.lower().split("@")[-1].endswith("g12.br"):
        return ""
    local = email.split("@")[0].replace("_", ".")
    return " ".join(p.capitalize() for p in local.split(".") if p)


def para_data(s):
    s = str(s or "").strip()
    for fmt in ("%Y-%m-%dT%H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d"):
        try:
            return datetime.strptime(s[:19], fmt)
        except ValueError:
            pass
    return s


def valor_txt(v):
    try:
        return f"{float(str(v).replace(',', '.')):.2f}"
    except ValueError:
        return str(v)


def ler_json(caminho):
    """Aceita uma resposta do conector, uma lista de respostas (uma por página) ou {"linhas": [...]}."""
    dados = json.load(open(caminho, encoding="utf-8"))
    paginas = dados if isinstance(dados, list) else [dados]
    linhas, internos = [], 0
    for pg in paginas:
        cols = pg.get("colunas") or ["Pedido", "Data", "Status", "Pagamento", "Email", "Aluno",
                                    "Responsável", "Produto", "ID produto", "Qtde", "Valor"]
        idx = {c: i for i, c in enumerate(cols)}
        internos += int(pg.get("pedidos_internos_ignorados") or 0)
        for l in pg.get("linhas", []):
            g = lambda c: l[idx[c]] if c in idx and idx[c] < len(l) else ""
            email = str(g("Email")).lower()
            linhas.append({
                "Pedido": g("Pedido"),
                "Data": para_data(g("Data")),
                "Status": g("Status"),
                "Método de Pagamento": g("Pagamento"),
                "Email": email,
                "Aluno": g("Aluno") or extrair_nome_aluno(email),
                "Produto": g("Produto"),
                "Qtde": g("Qtde"),
                "Valor": valor_txt(g("Valor")),
            })
    return linhas, internos


def ler_csv(caminho, produto=None, ano=None):
    """CSV exportado do WooCommerce: uma linha por pedido, itens em 'Product Item N ...'."""
    with open(caminho, encoding="utf-8-sig") as f:
        sep = ";" if ";" in f.readline() else ","
    linhas, internos = [], 0
    with open(caminho, encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f, delimiter=sep):
            status = str(row.get("status", "")).lower().strip()
            if status in STATUS_FORA:
                continue
            email = str(row.get("customer_email", "")).lower().strip()
            if email in EMAILS_INTERNOS:
                internos += 1
                continue
            data = para_data(row.get("order_date", ""))
            if ano and (not isinstance(data, datetime) or data.year != ano):
                continue
            for i in range(1, 51):
                nome = str(row.get(f"Product Item {i} Name") or "").strip()
                if not nome:
                    continue
                if produto and produto.lower().strip() not in nome.lower():
                    continue
                linhas.append({
                    "Pedido": row.get("order_number") or row.get("order_id", ""),
                    "Data": data,
                    "Status": STATUS_PT.get(status, status),
                    "Método de Pagamento": row.get("payment_method_title", ""),
                    "Email": email,
                    "Aluno": extrair_nome_aluno(email),
                    "Produto": nome,
                    "Qtde": row.get(f"Product Item {i} Quantity", ""),
                    "Valor": valor_txt(row.get(f"Product Item {i} Total", "0")),
                })
    return linhas, internos


def salvar(linhas, saida):
    for l in linhas:
        try:
            l["Pedido"] = int(l["Pedido"])
        except (TypeError, ValueError):
            pass
        try:
            l["Qtde"] = int(l["Qtde"])
        except (TypeError, ValueError):
            pass
    linhas.sort(key=lambda l: (str(l["Data"]), str(l["Pedido"])))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Relatório"
    ws.append(COLUNAS)
    for c in ws[1]:
        c.font = Font(bold=True)
    for l in linhas:
        ws.append([l[c] for c in COLUNAS])
    for i, nome in enumerate(COLUNAS, 1):
        ws.column_dimensions[get_column_letter(i)].width = LARGURAS[nome]
    for cel in ws["B"][1:]:
        if isinstance(cel.value, datetime):
            cel.number_format = "dd/mm/yyyy hh:mm"
    ws.freeze_panes = "A2"

    os.makedirs(os.path.dirname(os.path.abspath(saida)), exist_ok=True)
    wb.save(saida)


def main():
    p = argparse.ArgumentParser(description="Relatório de Vendas WooCommerce - GDV")
    p.add_argument("--json", help="Arquivo com a(s) resposta(s) de relatorio_vendas do conector")
    p.add_argument("--csv", help="CSV exportado do WooCommerce")
    p.add_argument("--produto", help="Só para --csv: trecho do nome do produto")
    p.add_argument("--ano", type=int, help="Só para --csv: ano do pedido")
    p.add_argument("--output", required=True, help="Caminho do .xlsx de saída")
    a = p.parse_args()

    if bool(a.json) == bool(a.csv):
        sys.exit("Informe --json ou --csv (um dos dois).")
    linhas, internos = ler_json(a.json) if a.json else ler_csv(a.csv, a.produto, a.ano)
    if not linhas:
        print("Nenhum item encontrado com esses filtros.")
        sys.exit(0)

    salvar(linhas, a.output)
    pedidos = {l["Pedido"] for l in linhas}
    por_status = {}
    for l in linhas:
        por_status[l["Status"]] = por_status.get(l["Status"], 0) + 1
    total = sum(float(l["Valor"]) for l in linhas if str(l["Valor"]).replace(".", "", 1).isdigit())
    print(json.dumps({"arquivo": a.output, "itens": len(linhas), "pedidos": len(pedidos),
                      "por_status": por_status, "valor_total": f"{total:.2f}",
                      "pedidos_internos_ignorados": internos}, ensure_ascii=False, indent=1))


if __name__ == "__main__":
    main()
